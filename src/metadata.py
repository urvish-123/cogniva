"""
Celebrity metadata loading.

Sources (later ones override earlier ones):
1. metadata/celebrities.csv   column "name" plus any detail columns
2. metadata/celebrities.xlsx  first sheet, any columns; the name column is
                              "Celebrity" or "Name"

Run directly to see which celebrities have enough photos:
python -m src.metadata
"""

import csv
import re

from .config import (
    ALLOWED_EXTENSIONS,
    CLEAN_DATASET_DIR,
    COVERAGE_REPORT_CSV,
    METADATA_CSV,
    METADATA_XLSX,
    MIN_IMAGES_PER_IDENTITY,
    RAW_DATASET_DIR,
)

NAME_COLUMNS = ("celebrity", "name")


def name_key(name: str) -> str:
    """'P. V. Sindhu', 'P_V_Sindhu' and 'pv sindhu' all map to 'pvsindhu'."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def make_record(name, details):
    lower = {k.lower(): v for k, v in details.items()}
    return {
        "name": name,
        "category": lower.get("category", ""),
        "language": lower.get("language", ""),
        "description": lower.get("description", ""),
        "details": details,
    }


def load_csv_metadata():
    metadata = {}
    if not METADATA_CSV.exists():
        return metadata

    with METADATA_CSV.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            name = clean(row.get("name"))
            if not name:
                continue
            details = {
                k: clean(v) for k, v in row.items()
                if k and k != "name" and clean(v)
            }
            metadata[name_key(name)] = make_record(name, details)
    return metadata


def load_excel_metadata():
    metadata = {}
    if not METADATA_XLSX.exists():
        return metadata

    from openpyxl import load_workbook

    sheet = load_workbook(METADATA_XLSX, read_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(h) for h in next(rows)]
    name_index = next(
        (i for i, h in enumerate(headers) if h.lower() in NAME_COLUMNS), 0)

    for row in rows:
        name = clean(row[name_index]) if name_index < len(row) else ""
        if not name:
            continue
        details = {
            header: clean(value)
            for i, (header, value) in enumerate(zip(headers, row))
            if i != name_index and header and clean(value)
        }
        metadata[name_key(name)] = make_record(name, details)
    return metadata


def load_metadata():
    metadata = load_csv_metadata()
    metadata.update(load_excel_metadata())
    return metadata


def count_images(directory):
    if not directory.is_dir():
        return 0
    return sum(p.suffix.lower() in ALLOWED_EXTENSIONS for p in directory.iterdir())


def folders_by_key(directory):
    if not directory.exists():
        return {}
    return {name_key(p.name): p for p in directory.iterdir() if p.is_dir()}


# Order in which statuses are printed (most urgent first).
STATUSES = {
    "missing_photos": "No photos added yet",
    "too_few_photos": f"Fewer than {MIN_IMAGES_PER_IDENTITY} photos added",
    "not_cleaned": "Photos added - run: python -m src.preprocess",
    "too_few_clean": f"Fewer than {MIN_IMAGES_PER_IDENTITY} photos passed cleaning",
    "ready": "Ready for recognition",
}


def photo_status(raw, clean):
    if raw == 0:
        return "missing_photos"
    if raw < MIN_IMAGES_PER_IDENTITY:
        return "too_few_photos"
    if clean == 0:
        return "not_cleaned"
    if clean < MIN_IMAGES_PER_IDENTITY:
        return "too_few_clean"
    return "ready"


def coverage_report():
    metadata = load_excel_metadata() or load_csv_metadata()
    raw_folders = folders_by_key(RAW_DATASET_DIR)
    clean_folders = folders_by_key(CLEAN_DATASET_DIR)

    rows = []
    for key, record in metadata.items():
        raw_folder = raw_folders.get(key)
        raw = count_images(raw_folder) if raw_folder else 0
        clean = count_images(clean_folders[key]) if key in clean_folders else 0
        rows.append({
            "name": record["name"],
            "folder": raw_folder.name if raw_folder else "",
            "photos_added": raw,
            "photos_clean": clean,
            "status": photo_status(raw, clean),
        })

    with COVERAGE_REPORT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    print(f"{'Celebrity':<32}{'Added':>7}{'Clean':>7}  Status")
    print("-" * 64)
    for status in STATUSES:
        for r in rows:
            if r["status"] == status:
                print(f"{r['name']:<32}{r['photos_added']:>7}"
                      f"{r['photos_clean']:>7}  {status}")

    print("-" * 64)
    print(f"Total celebrities: {len(rows)}")
    for status, label in STATUSES.items():
        count = sum(r["status"] == status for r in rows)
        print(f"  {label}: {count}")
    print(f"Total photos added: {sum(r['photos_added'] for r in rows)}")
    print(f"Report saved: {COVERAGE_REPORT_CSV}")


if __name__ == "__main__":
    coverage_report()
